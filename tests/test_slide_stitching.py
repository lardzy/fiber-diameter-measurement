from dataclasses import replace
from pathlib import Path

import cv2
import numpy as np
import pytest
from PySide6.QtGui import QImage

from fdm.services.digital_slide_store import DigitalSlideManifest, DigitalSlideStore, DigitalSlideTile
from fdm.services.slide_layout import LayoutTile, PairRegistrationResult, SlideLayoutSnapshot, load_layout, save_layout, validate_source
from fdm.services.slide_registration import estimate_pair, register_slide, solve_layout
from fdm.services.slide_raster import SlideRasterSource


def image(array):
    rgb = cv2.cvtColor(array, cv2.COLOR_GRAY2RGB)
    return QImage(rgb.data, rgb.shape[1], rgb.shape[0], rgb.shape[1] * 3, QImage.Format.Format_RGB888).copy()


def pixels(img):
    img = img.convertToFormat(QImage.Format.Format_RGB888)
    return np.frombuffer(img.constBits(), np.uint8).reshape(img.height(), img.bytesPerLine())[:, :img.width() * 3].copy()


@pytest.fixture
def scene():
    rng = np.random.default_rng(817)
    return cv2.GaussianBlur(rng.integers(0, 256, (1200, 1800), np.uint8), (0, 0), 1)


@pytest.mark.parametrize("dx,dy,axis", [(500, 3, "x"), (502.4, -2.6, "x"), (2.3, 402.7, "y")])
def test_pair_recovers_known_translation(scene, dx, dy, axis):
    a = scene[50:562, 50:690]
    xx, yy = np.meshgrid(np.arange(640, dtype=np.float32) + 50 + dx,
                         np.arange(512, dtype=np.float32) + 50 + dy)
    b = cv2.remap(scene, xx, yy, cv2.INTER_LINEAR)
    result = estimate_pair(a, b, dx=512 if axis == "x" else 0, dy=0 if axis == "x" else 410, axis=axis)
    assert result.accepted, result
    assert np.hypot(result.dx - dx, result.dy - dy) < .5


@pytest.mark.parametrize("overlap", [0, .02, .05, .10])
def test_low_conservative_overlap_never_searches(scene, monkeypatch, overlap):
    monkeypatch.setattr(cv2, "SIFT_create", lambda **_: pytest.fail("ineligible pair entered registration"))
    result = estimate_pair(scene[:512, :640], scene[:512, 600:1240], dx=640 * (1-overlap), dy=0, axis="x")
    assert not result.accepted and result.reason == "overlap"


def test_blank_and_straight_parallel_fibres_are_rejected():
    for a in (np.full((512, 640), 128, np.uint8), np.tile((np.sin(np.arange(640) / 3) * 100 + 128).astype(np.uint8), (512, 1))):
        assert not estimate_pair(a, a, dx=512, dy=0, axis="x").accepted


def make_slide(path, scene, *, mismatch=False):
    store = DigitalSlideStore.create(path, DigitalSlideManifest(1, 1152, 512, 640, 512, [0, 100]))
    for col in range(2):
        for z in range(2):
            store.write_tile(DigitalSlideTile(z, col*512, 0, 640, 512,
                stage_x=col*1000, focus_z=100*z + (5 if mismatch and col else 0)),
                image(scene[20:532, 20+col*500:660+col*500]))
    store.close()


def test_same_plane_layout_sidecar_and_source_integrity(tmp_path, scene):
    path = tmp_path / "sample.fdmslide"
    make_slide(path, scene)
    original = path.read_bytes()
    layout = register_slide(path, checkpoint=tmp_path / "checkpoint.json")
    assert layout.accepted_count == 1
    assert layout.pairs[0].verified_focus == (0, 1)
    for z in (0, 1):
        tiles = [t for t in layout.tiles if t.z_index == z]
        assert abs(tiles[1].x - tiles[0].x - 500) < .5
    save_layout(tmp_path / "sample.fdmstitch", layout)
    restored = load_layout(tmp_path / "sample.fdmstitch")
    validate_source(path, restored)
    assert restored == layout and path.read_bytes() == original
    moved = tmp_path / "renamed.fdmslide"
    moved.write_bytes(original)
    validate_source(moved, restored)
    store = DigitalSlideStore(moved)
    store._connection().execute("UPDATE tiles SET focus_z=focus_z+1 WHERE id=1")
    store.close()
    with pytest.raises(ValueError, match="内容已变化"):
        validate_source(moved, restored)


def test_same_index_different_focus_is_rejected(tmp_path, scene):
    path = tmp_path / "different-z.fdmslide"
    make_slide(path, scene, mismatch=True)
    result = register_slide(path)
    assert result.accepted_count == 0
    assert result.pairs[0].reason == "focus_mismatch"
    assert all(t.x == t.nominal_x and t.y == t.nominal_y for t in result.tiles)


def test_checkpoint_resume_does_not_recompute_unchanged_pairs(tmp_path, scene, monkeypatch):
    path = tmp_path / "sample.fdmslide"
    make_slide(path, scene)
    checkpoint = tmp_path / "progress.json"
    first = register_slide(path, checkpoint=checkpoint)
    def verify_only(*args, **kwargs):
        assert kwargs.get("fixed"), "recomputed checkpoint candidate search"
        return estimate_pair(*args, **kwargs)
    monkeypatch.setattr("fdm.services.slide_registration.estimate_pair", verify_only)
    assert register_slide(path, checkpoint=checkpoint) == first


def simple_layout(*, verified=True):
    tiles = tuple(LayoutTile(i+1, str(i), 0, 0, x, y, x+.25, y+.5, 80, 80, i, 0)
                  for i, (x, y) in enumerate(((0, 0), (60, 0), (0, 60), (60, 60))))
    pairs = tuple(PairRegistrationResult(str(a), str(b), axis, verified, "verified" if verified else "registration", verified_focus=(0,) if verified else ())
                  for a, b, axis in ((0,1,"x"),(2,3,"x"),(0,2,"y"),(1,3,"y")))
    return SlideLayoutSnapshot("test", tiles, pairs, 141, 141).sealed()


def test_single_source_composition_order_and_roi_invariance():
    layout = simple_layout()
    images = {t.tile_id: image(np.full((80,80), t.tile_id*40, np.uint8)) for t in layout.tiles}
    source = SlideRasterSource(layout, lambda t: images[t.tile_id])
    full = source.read_region(0, (0,0,141,141))
    part = source.read_region(0, (55,50,40,40))
    assert np.array_equal(pixels(full.image).reshape(141,141,3)[50:90,55:95], pixels(part.image).reshape(40,40,3))
    assert np.array_equal(full.unverified_seams[50:90,55:95], part.unverified_seams)
    other = SlideRasterSource(replace(layout, tiles=tuple(reversed(layout.tiles))), lambda t: images[t.tile_id])
    assert np.array_equal(pixels(other.read_region(0,(0,0,141,141)).image), pixels(full.image))
    assert full.coverage[1:140,1:140].all()


def test_unverified_seam_is_a_barrier_and_mask_roi_is_stable():
    source = SlideRasterSource(simple_layout(verified=False), lambda _: QImage())
    region = source.read_region(0,(0,0,141,141),pixels=False)
    assert region.unverified_seams.any()
    part = source.read_region(0,(65,0,12,141),pixels=False)
    assert np.array_equal(part.unverified_seams, region.unverified_seams[:,65:77])


def test_store_uses_same_repaired_pixels_and_coverage(tmp_path, scene):
    path = tmp_path / "sample.fdmslide"
    make_slide(path, scene)
    layout = register_slide(path)
    store = DigitalSlideStore(path)
    store.open_read_only()
    try:
        store.set_stitch_layout(layout)
        region = store.raster_source.read_region(0,(470,0,100,100))
        assert np.array_equal(pixels(store.render_viewport(x=470,y=0,width=100,height=100,z_index=0)), pixels(region.image))
        assert np.array_equal(store.viewport_coverage_mask(x=470,y=0,width=100,height=100,z_index=0), region.coverage)
    finally:
        store.close()


@pytest.fixture(scope="module")
def qapp():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def repaired_document(layout=None):
    from fdm.models import ImageDocument
    layout = layout or simple_layout(verified=False)
    return ImageDocument("repaired", "sample.fdmslide", (layout.width, layout.height),
                         document_kind="digital_slide", metadata={"stitch_layout": layout.to_dict(), "digital_slide": {"focus_index": 0}})


def test_whole_geometry_quality_holes_edit_and_project_roundtrip(qapp):
    from fdm.geometry import Point, Line
    from fdm.models import Measurement, ProjectState
    from fdm.services.slide_measurement_quality import quality_label
    document = repaired_document()
    line = Measurement("line", document.id, None, "manual", line_px=Line(Point(20,20),Point(120,20)))
    document.add_measurement(line)
    assert quality_label(line) == "接缝未验证"
    # Endpoints are inside individual fields; the connecting segment crosses.
    line.replace_line_geometry(line_px=Line(Point(10,20),Point(30,20)), snapped_line_px=None)
    document.mark_measurement_geometry_changed()
    assert not quality_label(line)
    area = Measurement("area", document.id, None, "polygon_area", measurement_kind="area",
                       polygon_px=[Point(30,30),Point(100,30),Point(100,100),Point(30,100)])
    document.add_measurement(area)
    assert quality_label(area)
    payload = ProjectState("test", [document]).to_dict()
    assert payload["project_schema_version"] == payload["min_reader_version"] == 3
    assert "slide-stitch-layout/v1" in payload["required_features"]
    loaded = ProjectState.from_dict(payload).documents[0]
    assert loaded.measurements[1].source_context == area.source_context
    assert loaded.metadata["stitch_layout"] == document.metadata["stitch_layout"]
    assert ProjectState("test", []).to_dict()["project_schema_version"] == 2


@pytest.mark.parametrize("dpr", [1, 1.25, 1.5, 2])
def test_renderer_native_matches_measurement_source(tmp_path, scene, qapp, dpr):
    from fdm.services.digital_slide_renderer import DigitalSlideRenderer, DigitalSlideRenderRequest
    path = tmp_path / "render.fdmslide"
    make_slide(path, scene)
    layout = register_slide(path)
    store = DigitalSlideStore(path)
    store.open_read_only()
    store.set_stitch_layout(layout)
    renderer = DigitalSlideRenderer(path, store.read_manifest(), stitch_layout=layout,
        cache_root=tmp_path / "cache", disk_cache_bytes=0, result_callback=lambda _:None, failure_callback=lambda _:None)
    try:
        request = DigitalSlideRenderRequest(1,"native",(430,30,200,200),(200,200),0,dpr,force_lod=0)
        frame = renderer._render(store,request)
        assert frame is not None and frame.pixel_exact
        assert frame.sampling == "single-source-linear-v1"
        assert pixels(frame.image).tolist() == pixels(store.render_viewport(x=430,y=30,width=200,height=200,z_index=0)).tolist()
        assert renderer._fingerprint == layout.layout_id
        assert renderer._tile_fingerprint == layout.source_digest
    finally:
        renderer.close()
        store.close()


def test_source_seam_barrier_and_confirmation_guard(tmp_path, scene, qapp):
    from fdm.geometry import Point
    from fdm.services.segmentation_source import digital_slide_segmentation_snapshot
    from fdm.ui.canvas import DocumentCanvas
    path = tmp_path / "failed-seam.fdmslide"
    make_slide(path,scene)
    layout = register_slide(path)
    layout = replace(layout, pairs=tuple(replace(p,accepted=False,verified_focus=()) for p in layout.pairs)).sealed()
    store = DigitalSlideStore(path)
    store.open_read_only()
    store.set_stitch_layout(layout)
    document = repaired_document(layout)
    try:
        snapshot = digital_slide_segmentation_snapshot(document,store,origin_px=Point(430,0),width=200,height=200,focus_index=0)
        assert snapshot.layout_id == layout.layout_id
        assert snapshot.unverified_seams.any()
        assert not (snapshot.valid_coverage & snapshot.unverified_seams).any()
        assert not snapshot.unverified_seams.flags.writeable
        canvas = DocumentCanvas()
        canvas._magic_segment.primary_debug_payload = {"segmentation_source":{"seam_truncated":True}}
        assert canvas.commit_magic_segment_preview()["reason"] == "unverified_seam"
        assert canvas.take_magic_commit_snapshot() is None
        canvas._fiber_quick.debug_payload = {"segmentation_source":{"seam_truncated":True}}
        assert not canvas.commit_fiber_quick_preview()["committed"]
        canvas.deleteLater()
    finally:
        store.close()


def test_calibrated_physical_and_pixel_steps_are_coupled_and_profile_bound():
    from fdm.settings import AppSettings
    from fdm.services.slide_capture_geometry import calibration_signature, calibrated_capture_settings, calibrated_plan_coordinates
    settings = AppSettings(digital_slide_pixel_stride_mode="calibrated_overlap",digital_slide_overlap_percent=20)
    settings.digital_slide_xy_calibration = {"signature":calibration_signature(settings,(1600,1200)),
        "x":{"pixels_per_step":.32,"cross_per_step":.002,"reliable":True},
        "y":{"pixels_per_step":.24,"cross_per_step":-.002,"reliable":True}}
    frozen = calibrated_capture_settings(settings,(1600,1200))
    assert frozen.digital_slide_x_stage_step == frozen.digital_slide_y_stage_step == 4000
    assert frozen.digital_slide_x_pixel_stride == 1280 and frozen.digital_slide_y_pixel_stride == 960
    assert settings.digital_slide_x_stage_step == 5000  # no mutation of preferences
    plan = [{"col":c,"row":r} for r in range(2) for c in range(2)]
    assert calibrated_plan_coordinates(plan, frozen,(1600,1200)) == (2888,2168)
    assert all(item["global_x"] >= 0 and item["global_y"] >= 0 for item in plan)
    with pytest.raises(ValueError, match="不一致"):
        calibrated_capture_settings(settings,(800,600))
    loaded = AppSettings.from_dict(settings.to_dict())
    assert loaded.digital_slide_xy_calibration == settings.digital_slide_xy_calibration


def test_quality_sheet_preserves_original_template_cells_and_macro_parts(tmp_path, qapp):
    from openpyxl import Workbook, load_workbook
    from fdm.geometry import Point, Line
    from fdm.models import Measurement
    from fdm.settings import RawRecordTemplate
    from fdm.services.raw_record_export import write_raw_record_template
    from fdm.services.export_service import ExportService
    import zipfile
    template = tmp_path / "template.xlsx"
    book = Workbook()
    book.active["B4"] = "keep me"
    book.active["B5"] = "=2+3"
    book.save(template)
    document = repaired_document()
    document.add_measurement(Measurement("m",document.id,None,"manual",line_px=Line(Point(20,20),Point(120,20))))
    output = write_raw_record_template(RawRecordTemplate("test",str(template),[]),tmp_path/"out.xlsx",
        documents=[document], measurement_rows=ExportService().build_measurement_rows([document]))
    result = load_workbook(output)
    assert result["Sheet"]["B4"].value == "keep me" and result["Sheet"]["B5"].value == "=2+3"
    assert "拼接质量说明" in result.sheetnames
    assert "接缝未验证" in str(list(result["拼接质量说明"].values))
    with zipfile.ZipFile(template) as original, zipfile.ZipFile(output) as target:
        assert original.read("xl/worksheets/sheet1.xml") == target.read("xl/worksheets/sheet1.xml")


def test_spawn_worker_completes_and_reclaims_resources(tmp_path, scene, qapp):
    from fdm.ui.slide_stitching_worker import SlideStitchingController, StitchJob
    from time import monotonic, sleep
    path = tmp_path / "spawn.fdmslide"
    make_slide(path,scene)
    controller = SlideStitchingController()
    results, errors = [], []
    controller.finished.connect(lambda *args:results.append(args))
    controller.failed.connect(lambda *args:errors.append(args))
    job = StitchJob("one",str(path),str(tmp_path/"result.fdmstitch"),str(tmp_path/"checkpoint.json"))
    try:
        controller.submit(job)
        controller.submit(job)  # duplicate active final is not accepted twice
        deadline = monotonic()+15
        while not results and not errors and monotonic()<deadline:
            qapp.processEvents()
            sleep(.01)
        assert not errors and len(results)==1
        assert results[0][1].accepted_count==1
        assert (tmp_path/"versions"/f"{results[0][1].layout_id}.fdmstitch").is_file()
        assert controller._process is None
    finally:
        controller.shutdown()
        controller.deleteLater()


@pytest.mark.parametrize("accepted", [True, False])
def test_repaired_view_is_independent_and_reopens_with_same_source(tmp_path, scene, qapp, monkeypatch, accepted):
    from unittest.mock import patch
    from fdm.settings import AppSettings
    from fdm.ui.main_window import MainWindow
    from fdm.models import Measurement, ProjectState
    from fdm.geometry import Point, Line
    from fdm.services.slide_layout import local_result_path
    monkeypatch.setattr("fdm.settings.settings_directory", lambda: tmp_path / "settings")
    path = tmp_path / "independent.fdmslide"
    make_slide(path,scene)
    layout = register_slide(path)
    if not accepted:
        layout = solve_layout(tuple(replace(t,x=t.nominal_x,y=t.nominal_y) for t in layout.tiles),
            tuple(replace(p,accepted=False,verified_focus=(),reason="registration") for p in layout.pairs),layout.source_digest)
    save_layout(local_result_path(path),layout)
    with patch("fdm.ui.main_window.AppSettingsIO.load",return_value=AppSettings()):
        window = MainWindow()
    try:
        window._add_digital_slide_document_from_path(path, document=None)
        original = window.current_document()
        first_group = original.create_group(label="棉", color="#00A0A0")
        original.add_measurement(Measurement("old",original.id,None,"manual",line_px=Line(Point(20,20),Point(30,20))))
        old_geometry = original.measurements[0].to_dict()
        window._open_current_slide_stitch_layout()
        repaired = window.current_document()
        assert repaired.id != original.id
        assert not repaired.measurements and original.measurements[0].to_dict() == old_geometry
        assert all(g.image_id == repaired.id and not g.measurement_ids for g in repaired.fiber_groups)
        assert first_group.id not in {g.id for g in repaired.fiber_groups}
        assert repaired.active_group_id in {g.id for g in repaired.fiber_groups}
        assert window._slide_stores[repaired.id].stitch_layout == layout
        if not accepted:
            assert "拼接检查" in window._document_display_name(repaired)
            assert all(t.x==t.nominal_x and t.y==t.nominal_y for t in layout.tiles)
        requests, skipped, _ = window._prepare_image_load_requests([(str(path), None)])
        assert not requests and skipped == 1
        saved = window.project.to_dict()
        restored = ProjectState.from_dict(saved)
        assert len(restored.documents) == 2
        assert restored.documents[1].metadata["stitch_layout"]["layout_id"] == layout.layout_id
        from fdm.project_io import ProjectIO
        project_path = tmp_path / "repair-project.fdmproj"
        ProjectIO.save(window.project, project_path)
        on_disk = ProjectIO.load(project_path)
        assert on_disk.documents[1].metadata["stitch_layout"]["layout_id"] == layout.layout_id
        local_result_path(path).unlink()
        # The project embeds the immutable layout; it does not need the sidecar.
        with patch("fdm.ui.main_window.AppSettingsIO.load",return_value=AppSettings()):
            reopened = MainWindow()
        try:
            for doc in on_disk.documents:
                reopened._add_digital_slide_document_from_path(path, document=doc, interaction_path_override=path)
            assert reopened._slide_stores[repaired.id].stitch_layout == layout
            assert len(reopened.project.documents) == 2
        finally:
            monkeypatch.setattr(reopened,"_confirm_close_documents",lambda *_:True)
            reopened.close()
            reopened.deleteLater()
        window.show()
        from time import monotonic, sleep
        deadline = monotonic() + 5
        while not window.current_canvas().pixel_work_enabled() and monotonic() < deadline:
            qapp.processEvents()
            sleep(.01)
        assert window.current_canvas().pixel_work_enabled(), window.current_canvas().pixel_work_unavailable_reason()
        # Exercise painting after the real background native frame arrives.
        window.grab().save(str(tmp_path / "repair-ui.png"))
        from fdm.services.export_service import ExportImageRenderMode
        export_path = tmp_path / "repair-overlay.png"
        window._render_overlay_image(repaired, export_path, include_measurements=False,
            include_scale=False, render_mode=ExportImageRenderMode.CURRENT_VIEWPORT)
        exported = QImage(str(export_path))
        assert exported.text("fdm.stitch.layout") == layout.layout_id
        origin = window.current_canvas().viewport_origin()
        expected = window._slide_stores[repaired.id].render_viewport(x=round(origin.x),y=round(origin.y),
            width=640,height=512,z_index=window.current_canvas().focus_index())
        assert np.array_equal(pixels(exported)[:-32],pixels(expected)[:-32])
    finally:
        monkeypatch.setattr(window,"_confirm_close_documents",lambda *_:True)
        window.close()
        window.deleteLater()
        qapp.processEvents()


def test_stage_suggestion_recomputes_pixels(qapp):
    from types import SimpleNamespace
    from unittest.mock import patch
    from PySide6.QtWidgets import QMessageBox
    from fdm.services.digital_slide_calibration import DigitalSlideCalibrationEstimate
    from fdm.ui.digital_slide_calibration import DigitalSlideCalibrationDialog
    from fdm.settings import AppSettings
    estimate = DigitalSlideCalibrationEstimate("x",1000,5,.2,6400,.95,4,4,(1600,1200),(1600,1200))
    target = SimpleNamespace(_current_estimate=estimate,_settings=AppSettings(),
        _apply_stage_checkbox=SimpleNamespace(isChecked=lambda:True),accept=lambda:None)
    with patch.object(QMessageBox,"question",return_value=QMessageBox.StandardButton.Yes):
        DigitalSlideCalibrationDialog._apply_result(target)
    assert target._applied_values["digital_slide_x_stage_step"] == 6400
    assert target._applied_values["digital_slide_x_pixel_stride"] == 1280
    assert target._applied_values["digital_slide_xy_calibration"]["x"]["reliable"]


def test_cross_layer_conflict_fails_closed(tmp_path, scene):
    path = tmp_path / "layer-conflict.fdmslide"
    store = DigitalSlideStore.create(path,DigitalSlideManifest(1,1152,512,640,512,[0,100]))
    for col in range(2):
        for z in range(2):
            start = 20 + col * (500 if z == 0 else 520)
            store.write_tile(DigitalSlideTile(z,col*512,0,640,512,stage_x=col*1000,focus_z=z*100),image(scene[20:532,start:start+640]))
    store.close()
    result = register_slide(path)
    assert not result.accepted_count and result.pairs[0].reason == "focus_conflict"
    assert all(tile.x == tile.nominal_x for tile in result.tiles)


def test_source_pixel_mutation_rejects_reuse(tmp_path, scene):
    path = tmp_path/"mutated.fdmslide"
    make_slide(path,scene)
    layout = register_slide(path)
    store = DigitalSlideStore(path)
    store._connection().execute("UPDATE tiles SET image_png=zeroblob(length(image_png)) WHERE id=1")
    store._connection().commit()
    store.close()
    with pytest.raises(ValueError,match="内容已变化"):
        validate_source(path,layout)


def test_periodic_two_dimensional_texture_is_not_accepted():
    cell = np.zeros((16,16),np.uint8)
    cv2.circle(cell,(8,8),4,210,-1)
    frame = np.tile(cell,(32,40))
    result = estimate_pair(frame,frame,dx=512,dy=0,axis="x")
    assert not result.accepted, result


def test_graph_conflict_and_isolated_field_are_conservative():
    tiles = tuple(LayoutTile(i+1,str(i),0,0,x,y,x,y,100,100,i%2,i//2)
        for i,(x,y) in enumerate(((0,0),(80,0),(0,80),(80,80),(400,0))))
    pairs = (
        PairRegistrationResult("0","1","x",True,"verified",80,0,.95,verified_focus=(0,)),
        PairRegistrationResult("0","2","y",True,"verified",0,80,.95,verified_focus=(0,)),
        PairRegistrationResult("1","3","y",True,"verified",0,80,.95,verified_focus=(0,)),
        PairRegistrationResult("2","3","x",True,"verified",95,0,.4,verified_focus=(0,)),
    )
    layout = solve_layout(tiles,pairs,"test")
    assert any(not p.accepted for p in layout.pairs)
    assert all(p.residual <= 1.5 for p in layout.pairs if p.accepted)
    isolated = next(t for t in layout.tiles if t.fov_id == "4")
    offset_x = next(t.x-t.nominal_x for t in layout.tiles if t.fov_id=="0")
    assert abs(isolated.x-400) <= 1  # no links, nominal placement


def test_worker_without_standard_streams(tmp_path, scene):
    import os
    import subprocess
    import sys
    path = tmp_path/"consoleless.fdmslide"
    make_slide(path,scene)
    script = tmp_path/"consoleless.py"
    script.write_text('''import sys, multiprocessing
sys.stdout = None
sys.stderr = None
from fdm.ui.slide_stitching_worker import StitchJob, _run_job
from pathlib import Path
path = Path(sys.argv[1])
context = multiprocessing.get_context("spawn")
if __name__ == "__main__":
    multiprocessing.freeze_support()
    events = context.Queue(8)
    stop = context.Event()
    process = context.Process(target=_run_job,args=(StitchJob("x",str(path),str(path.with_suffix(".fdmstitch")),str(path.with_suffix(".checkpoint"))),stop,events))
    process.start()
    process.join(12)
    if process.is_alive():
        process.terminate()
        process.join(2)
        raise SystemExit(2)
    code = process.exitcode
    process.close()
    events.cancel_join_thread()
    events.close()
    raise SystemExit(code)
''')
    env = {**os.environ,"PYTHONPATH":str(Path(__file__).resolve().parents[1]/"src")}
    result = subprocess.run([sys.executable,str(script),str(path)],env=env,capture_output=True,timeout=15)
    assert result.returncode == 0, result.stderr.decode(errors="replace")
    assert load_layout(path.with_suffix(".fdmstitch")).accepted_count == 1


@pytest.mark.parametrize("crosses_seam", [False, True])
def test_segmentation_barrier_preserves_safe_hole_fill(qapp, crosses_seam):
    from unittest.mock import patch
    from fdm.geometry import Point
    from fdm.services.mask_region import mask_region
    from fdm.services.prompt_segmentation import PromptSegmentationResult
    from fdm.settings import MagicSegmentToolMode
    from fdm.ui.prompt_segmentation_worker import PromptSegmentationWorker, PromptSegmentationRequest
    raw = np.zeros((60,80),bool)
    raw[10:45,10:70 if crosses_seam else 30] = True
    raw[20:25,18:23] = False
    class FakeService:
        def predict_polygon(self, **_kwargs):
            return PromptSegmentationResult(mask=mask_region(raw), polygon_px=[], area_rings_px=[], area_px=float(raw.sum()), metadata={})
    worker = PromptSegmentationWorker()
    worker._services["edge_sam"] = FakeService()
    seams = np.zeros_like(raw)
    seams[:,39:42] = True
    results, errors = [], []
    worker.succeeded.connect(lambda _doc,_request,result:results.append(result))
    worker.failed.connect(lambda *args:errors.append(args))
    request = PromptSegmentationRequest("slide",image(np.ones(raw.shape,np.uint8)),"source",1,
        [Point(12,12)],[],MagicSegmentToolMode.STANDARD,"add","edge_sam",False,
        valid_coverage=~seams,unverified_seams=seams,fill_draft_holes=True)
    with patch("fdm.ui.prompt_segmentation_worker.resolve_interactive_segmentation_backend",return_value=("edge_sam",None)):
        worker.infer(request)
    assert not errors and len(results)==1
    result=results[0]
    mask=result.mask.to_full_mask()
    assert not mask[:,39:].any()
    assert result.metadata["seam_truncated"] == crosses_seam
    assert result.metadata["holes_processed"] == (not crosses_seam)
    if not crosses_seam:
        assert mask[20:25,18:23].all()
    worker.deleteLater()


def test_final_verification_reports_partial_focus(tmp_path, scene, monkeypatch):
    from fdm.services.slide_registration import TranslationEvidence
    path=tmp_path/"partial.fdmslide"
    make_slide(path,scene)
    calls=0
    def check(*args, **kwargs):
        nonlocal calls
        if kwargs.get("fixed"):
            calls += 1
            if calls == 2:
                return TranslationEvidence(False,"registration")
        return estimate_pair(*args,**kwargs)
    monkeypatch.setattr("fdm.services.slide_registration.estimate_pair",check)
    layout=register_slide(path)
    pair=layout.pairs[0]
    assert pair.accepted and pair.verified_focus==(0,) and pair.reason=="partial_focus"
    assert pair.evidence_count==1
    assert pair.layers[1].reason=="final_seam_failed"
    raster=SlideRasterSource(layout,lambda _:QImage())
    assert list(raster.unverified_regions(1)) and not list(raster.unverified_regions(0))


def test_checkpoint_cancellation_and_algorithm_upgrade(tmp_path, scene, monkeypatch):
    import json
    path=tmp_path/"resume.fdmslide"
    make_slide(path,scene)
    checkpoint=tmp_path/"checkpoint.json"
    cancelled=False
    calls=0
    def interrupt_after_one(*args,**kwargs):
        nonlocal calls,cancelled
        calls+=1
        result=estimate_pair(*args,**kwargs)
        cancelled=True
        return result
    monkeypatch.setattr("fdm.services.slide_registration.estimate_pair",interrupt_after_one)
    with pytest.raises(InterruptedError):
        register_slide(path,checkpoint=checkpoint,cancelled=lambda:cancelled)
    assert len(json.loads(checkpoint.read_text())["evidence"])==1
    monkeypatch.setattr("fdm.services.slide_registration.estimate_pair",estimate_pair)
    layout=register_slide(path,checkpoint=checkpoint)
    assert layout.accepted_count==1
    saved=json.loads(checkpoint.read_text());saved["algorithm"]="older-acceptance-rules"
    checkpoint.write_text(json.dumps(saved))
    candidates=[]
    def count(*args,**kwargs):
        if not kwargs.get("fixed"):
            candidates.append(1)
        return estimate_pair(*args,**kwargs)
    monkeypatch.setattr("fdm.services.slide_registration.estimate_pair",count)
    register_slide(path,checkpoint=checkpoint)
    assert len(candidates)==2


def test_cancel_process_then_resume_and_publish_failure_keeps_source(tmp_path, scene, qapp):
    from fdm.ui.slide_stitching_worker import SlideStitchingController, StitchJob
    from fdm.ui.slide_stitching_publication import _Publisher
    from PySide6.QtWidgets import QMainWindow
    from time import monotonic,sleep
    path=tmp_path/"cancel.fdmslide"
    make_slide(path,scene)
    before=path.read_bytes()
    controller=SlideStitchingController()
    results,errors=[],[]
    controller.finished.connect(lambda *args:results.append(args))
    controller.failed.connect(lambda *args:errors.append(args))
    job=StitchJob("sample",str(path),str(tmp_path/"result.fdmstitch"),str(tmp_path/"progress.json"))
    window=QMainWindow()
    publisher=_Publisher(window)
    try:
        controller.submit(job)
        controller.cancel(job.key)
        controller.submit(job)
        deadline=monotonic()+20
        while not results and not errors and monotonic()<deadline:
            qapp.processEvents();sleep(.01)
        assert not errors and len(results)==1
        obstacle=tmp_path/"unavailable-directory"
        obstacle.write_text("not a directory")
        publisher.enqueue(Path(job.result),obstacle/"target.fdmstitch")
        deadline=monotonic()+10
        while publisher._process is not None and monotonic()<deadline:
            qapp.processEvents();sleep(.01)
        assert publisher._process is None
        assert "已保留本机结果" in window.statusBar().currentMessage()
        assert Path(job.result).is_file() and path.read_bytes()==before
        assert not publisher._pending
    finally:
        controller.shutdown();publisher.shutdown();window.close()
        controller.deleteLater();window.deleteLater();qapp.processEvents()


def test_missing_decoded_field_never_becomes_valid_measurement_pixels():
    raster=SlideRasterSource(simple_layout(),lambda _:QImage())
    with pytest.raises(ValueError,match="原始视场"):
        raster.read_region(0,(20,20,20,20))


def test_calibration_rejects_camera_crop_change_even_at_same_saved_size():
    from fdm.settings import AppSettings
    from fdm.services.slide_capture_geometry import calibration_signature,calibrated_capture_settings
    settings=AppSettings(digital_slide_pixel_stride_mode="calibrated_overlap",digital_slide_overlap_percent=20)
    settings.digital_slide_xy_calibration={"signature":calibration_signature(settings,(1600,1200)),
        "capture_frame_size":[5280,3960],
        "x":{"pixels_per_step":.2,"reliable":True},"y":{"pixels_per_step":.2,"reliable":True}}
    assert calibrated_capture_settings(settings,(1600,1200),source_frame_size=(5280,3960)).digital_slide_x_pixel_stride==1280
    with pytest.raises(ValueError,match="原始分辨率"):
        calibrated_capture_settings(settings,(1600,1200),source_frame_size=(2640,1980))


def test_stalled_metadata_publisher_is_bounded(tmp_path, qapp, monkeypatch):
    import multiprocessing
    from time import sleep
    from types import SimpleNamespace
    from PySide6.QtWidgets import QMainWindow
    from fdm.ui.slide_stitching_publication import _Publisher
    spawn = multiprocessing.get_context("spawn")
    monkeypatch.setattr("fdm.ui.slide_stitching_publication.multiprocessing.get_context",
        lambda _: SimpleNamespace(Process=lambda **_kwargs:spawn.Process(target=sleep,args=(60,),daemon=True)))
    window = QMainWindow()
    publisher = _Publisher(window)
    local = tmp_path / "kept.fdmstitch"
    local.write_text("local result")
    try:
        publisher.enqueue(local,tmp_path/"blocked-network"/"copy.fdmstitch")
        publisher._started -= 31
        for _ in range(5):
            publisher._poll()
            if publisher._process is None:
                break
            sleep(.01)
        assert publisher._process is None
        assert local.read_text() == "local result"
        assert "已保留本机结果" in window.statusBar().currentMessage()
    finally:
        publisher.shutdown();window.close();window.deleteLater();qapp.processEvents()


def test_capture_pause_survives_later_tile_notifications(qapp, monkeypatch):
    from fdm.ui.slide_stitching_worker import SlideStitchingController, StitchJob
    controller = SlideStitchingController()
    monkeypatch.setattr(controller, "_launch_next", lambda: None)
    job = StitchJob("capture", "source", "result", "checkpoint", final=False)
    try:
        # Pausing also works before the first four-field notification.
        controller.pause((job.key,))
        controller.submit(job)
        controller.submit(replace(job, final=True))
        assert not controller._pending
        controller.submit(replace(job, key="new-capture"))
        assert list(controller._pending) == ["new-capture"]
        controller.resume(job.key)
        controller.submit(replace(job, final=True))
        assert controller._pending[job.key].final
    finally:
        controller.shutdown();controller.deleteLater()


def test_z_approach_is_same_direction_and_obeys_soft_limit(tmp_path, qapp, monkeypatch):
    from unittest.mock import patch
    from fdm.settings import AppSettings
    from fdm.ui.main_window import MainWindow
    from fdm.services.motion_control import AXIS_X, AXIS_Y, AXIS_Z
    settings=AppSettings(digital_slide_z_backlash_steps=100,digital_slide_z_soft_limit=2000)
    with patch("fdm.ui.main_window.AppSettingsIO.load",return_value=settings):
        window=MainWindow()
    store=DigitalSlideStore.create(tmp_path/"z.fdmslide",DigitalSlideManifest(1,16,12,16,12,[-1000,0]))
    moves=[]
    def move(axis,position,**_kwargs):
        moves.append((axis,position));window._slide_motion.relative_pos[axis]=position;return True
    monkeypatch.setattr(window._slide_motion,"move_to",move)
    window._slide_motion.relative_pos={AXIS_X:0,AXIS_Y:0,AXIS_Z:500}
    try:
        window._slide_acquisition_store=store
        window._slide_acquisition_plan=[dict(z_index=0,global_x=0,global_y=0,stage_x=100,stage_y=0,focus_z=-1000,row=0,col=0)]
        window._slide_acquisition_index=0
        window._schedule_next_digital_slide_move()
        window._slide_acquisition_timer.stop()
        assert moves==[(AXIS_Z,-1100)] and window._slide_acquisition_timer_phase=="z_approach"
        window._on_slide_acquisition_timer_timeout()
        window._slide_acquisition_timer.stop()
        assert moves==[(AXIS_Z,-1100),(AXIS_X,100),(AXIS_Z,-1000)]
        assert window._slide_acquisition_timer_phase=="settle"
        window._slide_z_preapproached_index=-1
        window._slide_acquisition_plan[0]["focus_z"]=-1950
        failures=[]
        monkeypatch.setattr(window,"_fail_digital_slide_acquisition",failures.append)
        moves.clear()
        window._schedule_next_digital_slide_move()
        window._slide_acquisition_timer.stop()
        assert not moves and failures and "软限位" in failures[0]
    finally:
        window._slide_acquisition_store=None
        window._slide_acquisition_plan=[]
        store.close()
        monkeypatch.setattr(window,"_confirm_close_documents",lambda *_:True)
        window.close();window.deleteLater();qapp.processEvents()
