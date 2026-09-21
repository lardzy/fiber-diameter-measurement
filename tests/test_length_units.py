"""Calibration units remain consistent through input, persistence and export."""

import csv
import json

import pytest
from openpyxl import load_workbook

from fdm.geometry import Line, Point
from fdm.models import Calibration, CalibrationPreset, ImageDocument, Measurement, ProjectState
from fdm.project_io import ProjectIO
from fdm.services.export_service import ExportSelection, ExportService, SHEET_MEASUREMENT_DETAILS
from fdm.services.measurement_statistics import MeasurementMetric, MeasurementStatisticsService
from fdm.settings import AppSettings, AppSettingsIO
from fdm.ui.dialogs import CalibrationInputDialog, CalibrationPresetDialog
from fdm.ui.rendering import resolve_scale_overlay_value
from fdm.units import millimeters_per_unit


UNITS = ("nm", "um", "mm", "cm", "m")


def measured_document(unit: str, *, kind: str = "image") -> ImageDocument:
    document = ImageDocument(
        id=f"image-{unit}", path=f"sample-{unit}.png", image_size=(200, 100),
        document_kind=kind,
        calibration=Calibration("image_scale", 2.0, unit, "ruler"),
    )
    document.initialize_runtime_state()
    document.metadata["calibration_line"] = Line(Point(0, 0), Point(100, 0)).to_dict()
    document.add_measurement(Measurement(
        id=f"line-{unit}", image_id=document.id, fiber_group_id=None, mode="manual",
        line_px=Line(Point(0, 0), Point(6, 8)),
    ))
    document.add_measurement(Measurement(
        id=f"polyline-{unit}", image_id=document.id, fiber_group_id=None,
        mode="manual", measurement_kind="polyline",
        polyline_px=[Point(0, 0), Point(6, 0), Point(6, 8)],
    ))
    document.add_measurement(Measurement(
        id=f"area-{unit}", image_id=document.id, fiber_group_id=None,
        mode="manual", measurement_kind="area",
        polygon_px=[Point(0, 0), Point(10, 0), Point(10, 10), Point(0, 10)],
    ))
    return document


def test_calibration_dialogs_offer_five_units_and_keep_micrometer_default():
    for dialog in (CalibrationInputDialog(), CalibrationPresetDialog()):
        try:
            combo = dialog._unit_combo
            assert [combo.itemData(i) for i in range(combo.count())] == list(UNITS)
            assert [combo.itemText(i) for i in range(combo.count())] == ["nm", "μm", "mm", "cm", "m"]
            assert combo.currentData() == "um"
            for unit in UNITS:
                combo.setCurrentIndex(combo.findData(unit))
                if isinstance(dialog, CalibrationInputDialog):
                    assert dialog.values() == (100.0, unit, True)
                else:
                    assert dialog.values() == ("", 100.0, 10.0, 10.0, unit)
                    assert dialog._computed_label.text() == f"10.000000 px/{combo.currentText()}"
        finally:
            dialog.close()


@pytest.mark.parametrize("unit", [*UNITS, "µm", "μm", "legacy-custom"])
def test_editing_preset_preserves_its_unit_spelling(unit):
    dialog = CalibrationPresetDialog(
        initial_name="SEM", initial_pixel_distance=250.0,
        initial_actual_distance=20.0, initial_unit=unit,
    )
    try:
        dialog._name_edit.setText("SEM renamed")
        assert dialog.values() == ("SEM renamed", 250.0, 20.0, 12.5, unit)
    finally:
        dialog.close()


@pytest.mark.parametrize("unit,reference_length,expected_ppu", [
    ("nm", 10_000.0, 0.01),
    ("um", 10.0, 10.0),
    ("mm", 0.01, 10_000.0),
    ("cm", 0.001, 100_000.0),
    ("m", 0.00001, 10_000_000.0),
])
def test_same_physical_reference_in_every_unit(unit, reference_length, expected_ppu):
    dialog = CalibrationPresetDialog(
        initial_name="100 nm per pixel", initial_pixel_distance=100,
        initial_actual_distance=reference_length, initial_unit=unit,
    )
    try:
        name, pixels, actual, ppu, stored_unit = dialog.values()
    finally:
        dialog.close()
    assert ppu == pytest.approx(expected_ppu)
    preset = CalibrationPreset(name, ppu, stored_unit, pixels, actual, ppu)
    calibration = preset.to_calibration()
    factor = millimeters_per_unit(unit)
    assert calibration.px_to_unit(25) * factor == pytest.approx(0.0025)
    assert calibration.unit_to_px(0.0025 / factor) == pytest.approx(25)
    assert calibration.px_area_to_unit(625) * factor**2 == pytest.approx(6.25e-6)


@pytest.mark.parametrize("unit", [*UNITS, "µm", "μm"])
@pytest.mark.parametrize("kind", ["image", "digital_slide"])
def test_units_roundtrip_with_measurements_project_default_and_presets(tmp_path, unit, kind):
    document = measured_document(unit, kind=kind)
    preset = CalibrationPreset("preset", 2.0, unit, 100.0, 50.0, 2.0)
    settings_path = tmp_path / "settings.json"
    AppSettingsIO.save(AppSettings(calibration_presets=[preset]), settings_path)
    settings = AppSettingsIO.load(settings_path)
    assert settings.calibration_presets[0].to_dict() == preset.to_dict()

    project = ProjectState(
        version="test", documents=[document],
        project_default_calibration=document.calibration.as_project_default(),
    )
    path = ProjectIO.save(project, tmp_path / "units.fdmproj")
    loaded = ProjectIO.load(path)
    assert not loaded.load_issues
    assert loaded.project_default_calibration.unit == unit
    assert loaded.project_default_calibration.pixels_per_unit == 2.0
    result = loaded.documents[0]
    assert result.document_kind == kind
    assert result.calibration.to_dict() == document.calibration.to_dict()
    assert [m.display_value() for m in result.measurements] == [5.0, 7.0, 25.0]
    assert [m.display_label(result.calibration) for m in result.measurements] == [
        f"5.0000 {unit}", f"7.0000 {unit}", f"25.0000 {unit}²",
    ]
    assert resolve_scale_overlay_value(
        result, AppSettings(scale_overlay_length_value=50.0), image_to_output_scale=0.5,
    ) == (50.0, unit, 50.0)


def test_mixed_units_survive_csv_excel_scale_export_and_stay_separate_in_statistics(tmp_path):
    documents = [measured_document(unit) for unit in UNITS]
    project = ProjectState(version="test", documents=documents)
    outputs = ExportService().export_project(
        project, tmp_path,
        selection=ExportSelection(include_csv=True, include_excel=True, include_scale_json=True),
    )
    expected_units = [symbol for unit in UNITS for symbol in (unit, unit, f"{unit}²")]
    expected_values = [5.0, 7.0, 25.0] * len(UNITS)
    with outputs["measurement_details_csv"].open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert [row["单位"] for row in rows] == expected_units
    assert [float(row["结果"]) for row in rows] == expected_values
    workbook = load_workbook(outputs["xlsx"], read_only=True, data_only=True)
    try:
        values = list(workbook[SHEET_MEASUREMENT_DETAILS].values)
        unit_column = values[0].index("单位")
        value_column = values[0].index("结果")
        assert [row[unit_column] for row in values[1:]] == expected_units
        assert [row[value_column] for row in values[1:]] == expected_values
    finally:
        workbook.close()
    scales = [json.loads(path.read_text(encoding="utf-8")) for path in outputs["scale_jsons"]]
    assert [item["calibration"]["unit"] for item in scales] == list(UNITS)
    assert all(item["calibration"]["pixels_per_unit"] == 2 for item in scales)
    service = MeasurementStatisticsService()
    for metric, suffix, count, mean in (
        (MeasurementMetric.LENGTH, "", 2, 6.0), (MeasurementMetric.AREA, "²", 1, 25.0),
    ):
        statistics = service.summarize_documents(documents, metric=metric)
        assert [snapshot.unit for snapshot in statistics] == [f"{unit}{suffix}" for unit in UNITS]
        assert all(snapshot.mean == mean and snapshot.valid_count == count for snapshot in statistics)
