from copy import deepcopy
import csv
import json
import math

import pytest

from fdm.live_experiment import (ExperimentSession, MaterialSnapshot, default_materials, export_csv,
                                  export_xlsx, load_session, save_session, summarize)


def session():
    s = ExperimentSession()
    s.add_operator("甲")
    s.set_materials(default_materials(), direction="warp")
    s.set_materials(default_materials(), direction="weft")
    s.set_weights(70, 30)
    return s


def category(s, material="wool", direction="warp"):
    return next(row for row in summarize(s)["operators"][s.active_operator_id][direction]["categories"] if row["id"] == material)


def test_add_edit_delete_count_and_true_undo_redo():
    s = session()
    one = s.add_measurement(15, "wool")
    two = s.add_measurement(30, "wool")
    s.adjust_count("wool", 2)
    assert category(s)["mean"] == 22.5
    assert category(s)["total_count"] == 4
    assert s.undo()
    assert category(s)["extra_count"] == 0
    assert s.redo()
    assert category(s)["extra_count"] == 2
    s.edit_measurement(two.id, diameter_um=18, material_id="cashmere")
    assert category(s)["mean"] == 15
    assert category(s, "cashmere")["mean"] == 18
    s.undo()
    assert category(s)["mean"] == 22.5
    assert category(s, "cashmere")["measured_count"] == 0
    s.redo()
    s.delete_measurement(one.id)
    assert category(s)["measured_count"] == 0
    assert category(s)["extra_count"] == 2
    s.undo()
    assert s.measurements[one.id].diameter_um == 15
    s.adjust_count("wool", -2)
    assert not s.can_redo
    with pytest.raises(ValueError):
        s.adjust_count("wool", -1)
    assert len(s.measurements) == 2


def test_history_mutations_incrementally_match_full_rebuild():
    s = session()
    records = [s.add_measurement(10+i/7, "wool" if i%2 else "cashmere") for i in range(150)]
    for record in records[::7]:
        s.edit_measurement(record.id, diameter_um=record.diameter_um+4, material_id="cashmere")
    for record in records[::11]:
        s.delete_measurement(record.id)
    s.adjust_count("wool", 5)
    for _ in range(4):
        s.undo()
    for _ in range(2):
        s.redo()
    before = summarize(s)
    loaded = ExperimentSession.from_dict(s.to_dict())
    after = summarize(loaded)
    for key in ("warp", "weft"):
        for a, b in zip(before["operators"][s.active_operator_id][key]["categories"], after["operators"][s.active_operator_id][key]["categories"]):
            for field in ("mean", "stddev", "calculation_mean", "calculation_stddev", "content"):
                if a[field] is None:
                    assert b[field] is None
                else:
                    assert a[field] == pytest.approx(b[field], abs=1e-10, rel=1e-12)
            assert a["total_count"] == b["total_count"]


def test_configuration_undo_rebuilds_group_moments():
    s = session()
    s.add_measurement(15, "wool"); s.add_measurement(30, "wool")
    assert category(s)["calculation_mean"] == 23.75
    s.configure(group_width=2)
    assert category(s)["calculation_mean"] == 23
    s.undo()
    assert category(s)["calculation_mean"] == 23.75
    s.redo()
    assert category(s)["calculation_mean"] == 23


def test_limits_and_reminders_have_correct_boundary_semantics():
    m = MaterialSnapshot("x", "X", 1, lower_limit=10, upper_limit=20)
    assert [m.range_status(d) for d in (9, 10, 20, 21)] == ["below", "within", "within", "above"]
    s = session()
    s.configure(measured_reminder=2, total_reminder=3)
    s.add_measurement(15, "wool")
    s.add_measurement(18, "wool")
    assert s.reminder_state()["measured_reached"] == ["wool"]
    assert not s.reminder_state()["total_reached"]
    s.adjust_count("wool", 2)
    assert s.reminder_state()["total_reached"]  # no skipped-equality or N+1 bug
    assert s.reminder_state()["measured"]["wool"] == 2


def test_independent_directions_and_operator_definitions():
    s = session()
    a = s.active_operator_id
    s.add_measurement(15, "wool", direction="warp")
    s.add_measurement(30, "wool", direction="weft")
    b = s.add_operator("乙")
    assert s.operators[b].materials == s.operators[a].materials
    assert category(s)["measured_count"] == 0
    assert s.add_operator("甲") == a
    assert category(s)["mean"] == 15
    assert category(s, direction="weft")["mean"] == 30
    s.add_operator("丙")
    with pytest.raises(ValueError):
        s.add_operator("丁")
    s.active_operator_id = a
    with pytest.raises(ValueError):
        s.set_materials([default_materials()[1]], direction="warp")


@pytest.mark.parametrize("density,expected", [(None, "not_ready"), (0, "not_ready"), (1.3, "ready")])
def test_missing_density_is_not_a_zero_mass(density, expected):
    s = ExperimentSession()
    op = s.add_operator("A")
    s.set_materials([MaterialSnapshot("a", "A", density)])
    s.add_measurement(20, "a")
    result = summarize(s)["operators"][op]["warp"]
    assert result["status"] == expected
    assert result["categories"][0]["content"] == (100 if expected == "ready" else None)


def test_count_only_zero_weights_and_unsupported_material_are_explicit():
    s = session()
    s.adjust_count("wool", 1)
    result = summarize(s)
    assert result["fiber"]["status"] == "not_ready"
    assert category(s)["content"] is None
    s.add_measurement(20, "wool")
    s.set_weights(0, 0)
    assert summarize(s)["fabric"]["status"] == "not_ready"
    s.configure(method="FZ/T 32004-2009")
    assert summarize(s)["fiber"]["status"] == "invalid"
    s.configure(grouped=False)
    assert "grouping_required" in [r["code"] for r in summarize(s)["fiber"]["reasons"]]


@pytest.mark.parametrize("value", [float("nan"), float("inf"), -1, 0, True, 1e308])
def test_invalid_measurement_cannot_change_session(value):
    s = session()
    revision = s.revision
    with pytest.raises(ValueError):
        s.add_measurement(value, "wool")
    assert s.revision == revision
    assert not s.measurements


def test_net_zero_count_history_retains_material_definition_for_export():
    s = session()
    s.adjust_count("wool", 1)
    s.adjust_count("wool", -1)
    with pytest.raises(ValueError):
        s.set_materials([default_materials()[1]])
    assert len(s.count_adjustments) == 2


def test_summary_does_not_scan_records_or_expose_cache():
    class NoScan(dict):
        def values(self):
            raise AssertionError("high-frequency summary scanned measurement history")
        def items(self):
            raise AssertionError("high-frequency summary scanned measurement history")
        def __iter__(self):
            raise AssertionError("high-frequency summary scanned measurement history")
    s = session()
    s.add_measurement(20, "wool")
    s.measurements = NoScan(s.measurements)
    result = summarize(s)
    result["fiber"]["categories"][0]["content"] = -100
    assert summarize(s)["fiber"]["categories"][0]["content"] == 100
    s.add_measurement(30, "wool")
    assert category(s)["mean"] == 25


def test_async_evidence_is_not_a_measurement_undo_and_survives_redo():
    s = session()
    evidence = {"state": "pending", "asset_id": "asset1"}
    provenance = {"calibration": {"scale": 0.25}, "frame_id": 42}
    record = s.add_measurement(15, "wool", evidence=evidence, provenance=provenance,
                               endpoints=((1, 2), (3, 4)), frame_id=42)
    evidence["state"] = "modified"
    provenance["calibration"]["scale"] = 900
    assert record.evidence["state"] == "pending"
    assert record.provenance["calibration"]["scale"] == 0.25
    s.undo()
    assert not s.measurements
    assert s.set_evidence(record.id, {"state": "saved", "path": "evidence/x.png", "sha256": "abc"})
    s.redo()
    assert s.measurements[record.id].evidence["state"] == "saved"
    s.undo()
    assert not s.measurements
    assert not s.set_evidence("unknown", {"state": "saved"})


def test_json_roundtrip_and_atomic_failure_preserves_file(tmp_path):
    s = session()
    s.add_measurement(20, "wool", evidence={"state": "failed", "reason": "磁盘已满"})
    s.adjust_count("wool", 2); s.adjust_count("wool", -1)
    s.fields["field1"] = {"state": "closed", "reason": "movement"}
    path = tmp_path / "试验.json"
    save_session(s, path)
    restored = load_session(path)
    assert restored.to_dict() == s.to_dict()
    assert restored.can_undo is False  # editing history is session-local, measurements persist
    original = path.read_bytes()
    s.fields["bad"] = {"value": float("nan")}
    with pytest.raises(ValueError):
        save_session(s, path)
    assert path.read_bytes() == original
    bad = json.loads(original)
    bad["schema_version"] = 999
    with pytest.raises(ValueError):
        ExperimentSession.from_dict(bad)
    bad = json.loads(original)
    bad["measurements"][0]["operator_id"] = "missing"
    with pytest.raises(ValueError):
        ExperimentSession.from_dict(bad)


def test_csv_xlsx_export_keeps_count_separate_and_matches_summary(tmp_path):
    from openpyxl import load_workbook
    s = session()
    s.sample_number = "样品编号甲"
    s.add_measurement(15, "wool", frame_id=3, field_id="f", evidence={"path": "evidence/3.png", "state": "saved"})
    s.add_measurement(30, "wool")
    s.add_measurement(18, "cashmere")
    s.adjust_count("wool", 1)
    csv_path = export_csv(s, tmp_path / "records.csv")
    with csv_path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    assert len(rows) == 5
    assert rows[-1][0] == "额外计数"
    assert rows[-1][7] == ""
    assert rows[-1][8] == "1"
    path = export_xlsx(s, tmp_path / "report.xlsx")
    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == ["试验说明", "原始记录", "单人单向统计", "含量汇总", "材料快照"]
    assert workbook["单人单向统计"].cell(2, 13).value == pytest.approx(85.33239059467755)
    assert workbook["原始记录"].cell(5, 8).value is None
    assert workbook["原始记录"].cell(5, 9).value == 1
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()


def test_metadata_edit_undo_and_batch_atomicity():
    s = session()
    s.configure(sample_name="新样品", sample_number="S-1", name="试验一")
    assert s.sample_name == "新样品"
    s.undo()
    assert s.sample_name == ""
    s.redo()
    assert s.sample_number == "S-1"
    with pytest.raises(ValueError):
        s.configure(sample_name=5)
    before = s.to_dict()
    with pytest.raises(ValueError):
        s.add_measurements_batch([{"diameter_um": 15, "material_id": "wool"},
                                  {"diameter_um": float("nan"), "material_id": "wool"}])
    assert s.to_dict() == before
    with pytest.raises(ValueError):
        s.add_measurements_batch([{"diameter_um": 15, "material_id": "wool", "id": "same"},
                                  {"diameter_um": 18, "material_id": "cashmere", "id": "same"}])
    assert s.to_dict() == before
    revision = s.revision
    records = s.add_measurements_batch([{"diameter_um": 15, "material_id": "wool"},
                                        {"diameter_um": 18, "material_id": "cashmere"}],
                                       provenance={"algorithm_revision": "fixture", "source_frame_id": 42})
    assert s.revision == revision + 1
    assert len(s.measurements) == 2
    batch = next(iter(s.batches.values()))
    assert batch["record_ids"] == [record.id for record in records]
    assert all(record.provenance["batch_id"] == batch["id"] for record in records)
    assert ExperimentSession.from_dict(s.to_dict()).to_dict() == s.to_dict()
    s.undo()
    assert not s.measurements and not s.batches
    s.redo()
    assert len(s.measurements) == 2 and len(s.batches) == 1


def test_accumulator_overflow_is_rejected_before_any_ledger_mutation():
    s = session()
    s.add_measurement(1e154, "wool")
    before = s.to_dict()
    with pytest.raises(ValueError):
        s.add_measurement(1e154, "wool")
    assert s.to_dict() == before
    s = session()
    before = s.to_dict()
    with pytest.raises(ValueError):
        s.add_measurements_batch([{"diameter_um": 1e154, "material_id": "wool"}] * 2)
    assert s.to_dict() == before


def test_builtin_catalog_retains_recovered_density_bounds_and_versions():
    from fdm.live_experiment import built_in_materials
    normal, all_types = built_in_materials("normal"), built_in_materials()
    assert len(normal) == 37 and len(all_types) == 40
    assert len({material.id for material in all_types}) == 40
    assert [(m.id, m.name, m.density) for m in normal] == [(m.id, m.name, m.density) for m in all_types[:37]]
    by_name = {m.name: m for m in all_types}
    expected = {"棉": (1.54, "cotton"), "亚麻": (1.5, "flax"), "大麻": (1.48, "hemp"),
                "苎麻": (1.51, "ramie"), "罗布麻": (1.5, "apocynum"),
                "桑蚕丝": (1.36, "other"), "锦纶": (1.14, "other"), "莱赛尔纤维": (1.52, "other")}
    for name, value in expected.items():
        assert (by_name[name].density, by_name[name].family) == value
    assert by_name["山羊绒"].upper_limit == 30
    assert by_name["山羊毛"].lower_limit == 30
    assert by_name["山羊绒"].upper_suggestion_id == "goat_hair"
    assert by_name["牦牛绒"].upper_limit == 35
    assert by_name["驼绒"].upper_limit == 40
    assert by_name["绵羊毛"].lower_limit is None
    assert by_name["绵羊毛"].provenance["source_record_offset"] == 0
    for name in ("棕色棉", "绿色棉", "白棉"):
        assert by_name[name].density == 0 and by_name[name].family == "cotton"
        assert by_name[name].provenance["source_versions"] == ["cu6mt"]
        s = ExperimentSession()
        s.add_operator("A")
        s.set_materials([by_name[name]])
        s.add_measurement(20, by_name[name].id)
        assert summarize(s)["status"] == "not_ready"
    all_types[0].provenance["source_versions"].clear()
    assert built_in_materials()[0].provenance["source_versions"] == ["cu5", "cu6", "cu6mt"]


def test_workspace_options_persist_partial_updates_and_undo():
    s = ExperimentSession()
    assert s.workspace_options == {"gesture": "click", "mode": "content", "classification": "before"}
    options = {"gesture": "drag"}
    s.configure(workspace_options=options)
    options["gesture"] = "changed outside"
    assert s.workspace_options == {"gesture": "drag", "mode": "content", "classification": "before"}
    s.configure(workspace_options={"mode": "diameter"})
    assert s.workspace_options == {"gesture": "drag", "mode": "diameter", "classification": "before"}
    s.undo()
    assert s.workspace_options == {"gesture": "drag", "mode": "content", "classification": "before"}
    s.redo()
    payload = s.to_dict()
    loaded = ExperimentSession.from_dict(json.loads(json.dumps(payload)))
    assert loaded.workspace_options == {"gesture": "drag", "mode": "diameter", "classification": "before"}
    payload["workspace_options"]["mode"] = "modified"
    assert s.workspace_options["mode"] == "diameter"
    old = s.to_dict()
    old.pop("workspace_options")
    assert ExperimentSession.from_dict(old).workspace_options == {"gesture": "click", "mode": "content", "classification": "before"}


@pytest.mark.parametrize("options", [[], None, {"gesture": "double"}, {"mode": "other"}, {"classification": "sometimes"}, {"camera": 2}])
def test_invalid_workspace_options_leave_session_unchanged(options):
    s = ExperimentSession()
    before = s.to_dict()
    with pytest.raises(ValueError):
        s.configure(workspace_options=options)
    assert s.to_dict() == before


def test_unchanged_width_and_method_only_do_not_visit_record_collections():
    class NoTraversal(dict):
        def values(self):
            raise AssertionError("small configuration traversed historical records")
        def __iter__(self):
            raise AssertionError("small configuration traversed historical records")
    s = session()
    s.add_measurement(15, "wool")
    s.measurements = NoTraversal(s.measurements)
    s.count_adjustments = NoTraversal(s.count_adjustments)
    before = s.revision
    assert s.configure_deferred(group_width=2.5) is None
    assert s.revision == before
    assert s.configure_deferred(method="SN/T 1062-2002", group_width=2.5, measured_reminder=501) is None
    assert s.method == "SN/T 1062-2002" and s.measured_reminder == 501
    assert s.group_width == 2.5 and s.revision == before + 1


def test_deferred_group_width_snapshot_and_atomic_independent_moments():
    s = session()
    records = [s.add_measurement(value, "wool") for value in (15, 30)]
    s.adjust_count("wool", 3)
    before = s.to_dict()
    request = s.configure_deferred(group_width=2, method="SN/T 1062-2002")
    assert request is not None and request.measurements == tuple(records)
    assert request.measurements[0] is records[0]  # no deep copy of evidence/pixels
    assert s.to_dict() == before
    result = request.build()
    assert s.to_dict() == before
    s.apply_rebuild(request, result)
    row = category(s)
    # Independent arithmetic: 15 -> 15 and 30 -> 31 for 2 μm bins.
    assert row["mean"] == 22.5 and row["calculation_mean"] == 23
    assert row["calculation_stddev"] == pytest.approx(math.sqrt(128))
    assert row["extra_count"] == 3 and row["total_count"] == 5
    assert row["weight"] == pytest.approx(23 ** 2 * 5 * 1.31)
    assert s.method == "SN/T 1062-2002"


@pytest.mark.parametrize("mutation", ["append", "edit", "count", "configure", "evidence"])
def test_deferred_rebuild_rejects_changed_revision_without_overwriting(mutation):
    s = session()
    record = s.add_measurement(15, "wool")
    request = s.configure_deferred(group_width=2)
    if mutation == "append":
        s.add_measurement(22, "wool")
    elif mutation == "edit":
        s.edit_measurement(record.id, diameter_um=19)
    elif mutation == "count":
        s.adjust_count("wool", 1)
    elif mutation == "configure":
        s.configure(measured_reminder=444)
    else:
        s.set_evidence(record.id, {"state": "failed", "error": "disk full"})
    before = s.to_dict()
    undo = list(s._undo)
    with pytest.raises(ValueError, match="数据已变化"):
        s.apply_rebuild(request, request.build())
    assert s.to_dict() == before and s._undo == undo


def test_deferred_width_cancel_failure_and_history_keep_old_state():
    from concurrent.futures import CancelledError
    s = session()
    for _ in range(10):
        s.add_measurement(15, "wool")
    before = s.to_dict()
    request = s.configure_deferred(group_width=2)
    with pytest.raises(CancelledError):
        request.build(lambda: True)
    assert s.to_dict() == before
    overflow = s.configure_deferred(group_width=1e154)
    with pytest.raises((ValueError, OverflowError)):
        overflow.build()
    assert s.to_dict() == before
    s.apply_rebuild(request, request.build())
    for action, expected_width, expected_mean in (("undo", 2.5, 16.25), ("redo", 2.0, 15.0)):
        stack_before = (list(s._undo), list(s._redo))
        request = s.history_deferred(action)
        assert request is not None
        assert (s._undo, s._redo) == stack_before
        s.apply_rebuild(request, request.build())
        assert s.group_width == expected_width
        assert category(s)["calculation_mean"] == expected_mean


def test_custom_group_width_survives_json_and_deferred_roundtrip():
    s = session()
    s.add_measurement(15, "wool")
    request = s.configure_deferred(group_width=1.75)
    s.apply_rebuild(request, request.build())
    loaded = ExperimentSession.from_dict(json.loads(json.dumps(s.to_dict())))
    assert loaded.group_width == 1.75
    assert category(loaded)["calculation_mean"] == 14.875


def test_rebuild_cancellation_during_scan_never_publishes_partial_aggregates():
    from concurrent.futures import CancelledError
    s = session()
    for value in range(300):
        s.add_measurement(10 + value / 100, "wool")
    before = summarize(s)
    request = s.configure_deferred(group_width=2)
    checkpoints = []
    def cancel_after_first_chunk():
        checkpoints.append(True)
        return len(checkpoints) == 2
    with pytest.raises(CancelledError):
        request.build(cancel_after_first_chunk)
    assert len(checkpoints) == 2
    assert summarize(s) == before
