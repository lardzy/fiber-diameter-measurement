"""Snapshot exports; no live Excel/COM dependency on the measurement path."""
from __future__ import annotations

import csv
from io import StringIO
import json
from pathlib import Path
from typing import Any

from fdm.atomic_io import atomic_replace_file, atomic_write_text, staged_path_for
from .models import ExperimentSession
from .statistics import summarize


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(",", ":"))


def _measurement_rows(session: ExperimentSession) -> list[list[Any]]:
    rows: list[list[Any]] = [["记录类型", "记录 ID", "操作者 ID", "操作者", "方向", "材料 ID", "材料", "直径 μm",
                            "额外计数变化", "视场 ID", "帧 ID", "测量方法", "时间", "端点", "图像证据", "测量来源"]]
    for record in session.measurements.values():
        material = session.material(record.material_id, operator_id=record.operator_id, direction=record.direction)
        rows.append(["实测", record.id, record.operator_id, session.operators[record.operator_id].name,
                     record.direction, material.id, material.name, record.diameter_um, "", record.field_id,
                     record.frame_id, record.measurement_method, record.created_at, _json(record.endpoints),
                     _json(record.evidence), _json(record.provenance)])
    for event in session.count_adjustments.values():
        material = session.material(event.material_id, operator_id=event.operator_id, direction=event.direction)
        rows.append(["额外计数", event.id, event.operator_id, session.operators[event.operator_id].name,
                     event.direction, material.id, material.name, "", event.delta, event.field_id,
                     "", "", event.created_at, "", "", event.note])
    return rows


def export_csv(session: ExperimentSession, path: str | Path) -> Path:
    stream = StringIO(newline="")
    csv.writer(stream).writerows(_measurement_rows(session))
    # BOM supports the Windows Excel default UTF-8 opening path.
    return atomic_write_text(path, stream.getvalue(), encoding="utf-8-sig")


def export_xlsx(session: ExperimentSession, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    summary = summarize(session)
    workbook = Workbook()
    metadata = workbook.active
    metadata.title = "试验说明"
    for row in [["项目", "值"], ["试验 ID", session.id], ["试验名称", session.name],
                ["样品编号", session.sample_number], ["样品名称", session.sample_name],
                ["计算方法（软件规则）", session.method], ["计算规则版本", summary["calculation_revision"]],
                ["模板版本", session.profile], ["分组", session.grouped], ["组距 μm", session.group_width],
                ["状态", summary["status"]], ["标准符合性已验收", False],
                ["算法修订说明", "保留已恢复规则；修正固定行数/桶数截断、单样本空值及舍入边界"],
                ["注意", "实测、额外计数分别记录；含量未就绪时不填伪数值。图片为证据引用。"]]:
        metadata.append(row)
    measurements = workbook.create_sheet("原始记录")
    for row in _measurement_rows(session):
        measurements.append(row)
    category_sheet = workbook.create_sheet("单人单向统计")
    category_sheet.append(["操作者", "方向", "材料 ID", "材料", "密度", "实测数", "额外计数", "总数",
                           "实测均径", "实测标准差", "计算均径", "计算标准差", "含量 %", "状态", "原因"])
    for op in summary["operators"].values():
        for direction in ("warp", "weft"):
            result = op[direction]
            for row in result["categories"]:
                category_sheet.append([op["name"], direction, row["id"], row["name"], row["density"],
                    row["measured_count"], row["extra_count"], row["total_count"], row["mean"], row["stddev"],
                    row["calculation_mean"], row["calculation_stddev"], row["content"], result["status"],
                    "; ".join(reason["message"] for reason in result["reasons"])])
    report = workbook.create_sheet("含量汇总")
    report.append(["汇总层级", "操作者", "材料 ID", "材料", "未舍入含量 %", "报告含量 %", "状态", "原因"])
    for op in summary["operators"].values():
        result = op["mixed"]
        for row in result["categories"]:
            report.append(["单人经纬加权", op["name"], row["id"], row["name"], row["content"], row["content"],
                           result["status"], "; ".join(r["message"] for r in result["reasons"])])
    for key, label in (("fiber", "多人纤维汇总"), ("fabric", "多人织物汇总")):
        result = summary[key]
        if not result["categories"]:
            report.append([label, "", "", "", None, None, result["status"], "; ".join(r["message"] for r in result["reasons"])])
        for row in result["categories"]:
            report.append([label, "", row["id"], row["name"], row["raw_content"], row["content"], result["status"],
                           "; ".join(r["message"] for r in result["reasons"])])
    materials = workbook.create_sheet("材料快照")
    materials.append(["操作者", "方向", "材料 ID", "名称", "材料族", "密度", "直径下限", "直径上限", "定义来源"])
    for operator in session.operators.values():
        for direction, definitions in operator.materials.items():
            for material in definitions:
                materials.append([operator.name, direction, material.id, material.name, material.family, material.density,
                                  material.lower_limit, material.upper_limit, _json(material.provenance)])
    if session.batches:
        batches = workbook.create_sheet("批次来源")
        batches.append(["批次 ID", "创建时间", "全部原始记录 ID", "共同算法与图像来源"])
        for batch in session.batches.values():
            batches.append([batch["id"], batch["created_at"], _json(batch["record_ids"]), _json(batch["provenance"])])
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24536B")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(48, max(12, max(len(str(c.value or "")) for c in column) + 2))
            for cell in column:
                # Preserve literal sample/material names in spreadsheet cells.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    with staged_path_for(path, suffix=".xlsx") as staged:
        workbook.save(staged)
        workbook.close()
        return atomic_replace_file(staged, path)


session_export_xlsx = export_xlsx
