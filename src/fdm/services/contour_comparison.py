"""Native-pixel silhouette comparison in a fixed, explicitly defined axis frame.

This measures shape at the same physical height, not material strain.  Axis
length is never a scale factor.  Display contours and previews are disposable;
all measurements use boundaries of the native binary pixel cells.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass, replace
from io import BytesIO
import json
import math
import os
from pathlib import Path
import tempfile
from typing import Callable
import zipfile

import cv2
import numpy as np
from PIL import Image, ImageOps

from fdm.cancellation import CancellationToken

MAX_PIXELS = 36_000_000
MAX_SAMPLES = 5000
SCHEMA = "fdm.contour-comparison"


def _check(token: CancellationToken | None) -> None:
    if token is not None:
        token.raise_if_cancelled()


def _positive(value: float, label: str) -> float:
    value = float(value)
    if not math.isfinite(value) or value <= 0:
        raise ValueError(f"{label}必须为大于零的有限数值。")
    return value


@dataclass(frozen=True)
class ContourAxis:
    origin: tuple[float, float]
    direction_point: tuple[float, float]

    def __post_init__(self) -> None:
        points = np.asarray((self.origin, self.direction_point), dtype=float)
        if points.shape != (2, 2) or not np.isfinite(points).all():
            raise ValueError("中线需要两个有效坐标。")
        if np.linalg.norm(points[1] - points[0]) < 1e-6:
            raise ValueError("中线的两点不能重合。")
        object.__setattr__(self, "origin", tuple(map(float, points[0])))
        object.__setattr__(self, "direction_point", tuple(map(float, points[1])))

    def basis(self) -> np.ndarray:
        d = np.subtract(self.direction_point, self.origin)
        d /= np.linalg.norm(d)
        return np.array(((d[1], -d[0]), (d[0], d[1])))

    def to_world(self, points: np.ndarray, scale: float = 1.0) -> np.ndarray:
        return (np.asarray(points) - self.origin) @ self.basis().T * scale

    def to_image(self, points: np.ndarray, scale: float = 1.0) -> np.ndarray:
        return np.asarray(points) / scale @ self.basis() + self.origin


@dataclass(frozen=True, eq=False)
class ContourFrame:
    label: str
    rgba: np.ndarray
    mask: np.ndarray
    axis: ContourAxis
    mm_per_pixel: float | None = None
    warnings: tuple[str, ...] = ()
    source_path: str = ""
    edited: bool = False
    axis_confirmed: bool = False

    def __post_init__(self) -> None:
        if (self.rgba.ndim != 3 or self.rgba.shape[2] != 4
                or self.rgba.dtype != np.uint8):
            raise ValueError("输入图像需要 RGBA8 像素。")
        h, w = self.rgba.shape[:2]
        if min(h, w) < 2 or h * w > MAX_PIXELS:
            raise ValueError("图片尺寸无效，单图最多支持 3600 万像素。")
        if self.mask.shape != (h, w) or self.mask.dtype != np.bool_:
            raise ValueError("轮廓掩膜必须与原图尺寸一致。")
        self.rgba.setflags(write=False)
        self.mask.setflags(write=False)
        if self.mm_per_pixel is not None:
            object.__setattr__(self, "mm_per_pixel", _positive(self.mm_per_pixel, "像素比例"))

    @property
    def scale(self) -> float:
        return self.mm_per_pixel if self.mm_per_pixel is not None else 1.0


@dataclass(frozen=True)
class ContourSection:
    height: float
    before_intervals: tuple[tuple[float, float], ...]
    after_intervals: tuple[tuple[float, float], ...]
    left_before: float | None
    left_after: float | None
    right_before: float | None
    right_after: float | None
    left_change: float | None
    right_change: float | None
    span_before: float | None
    span_after: float | None
    span_change: float | None
    status: str
    # Ordered intersection endpoints, not tracked material points. No automatic
    # pairing across topology changes or across the centerline.
    boundary_changes: tuple[tuple[float | None, float | None], ...] = ()


@dataclass(frozen=True)
class ContourComparison:
    unit: str
    step: float
    sections: tuple[ContourSection, ...]
    # Bounds use (left, top, right, bottom) in the axis coordinate system.
    before_bounds: tuple[float, float, float, float]
    after_bounds: tuple[float, float, float, float]
    before_area: float
    after_area: float
    warnings: tuple[str, ...]

    @property
    def summary(self) -> dict[str, float]:
        b, a = self.before_bounds, self.after_bounds
        return {
            "左侧最外缘变化": b[0] - a[0],
            "右侧最外缘变化": a[2] - b[2],
            "上端向外变化": b[1] - a[1],
            "下端向外变化": a[3] - b[3],
            "纵向总长变化": (a[3] - a[1]) - (b[3] - b[1]),
            "投影面积变化": self.after_area - self.before_area,
            "投影面积变化率 (%)": 100 * (self.after_area / self.before_area - 1),
        }


def default_axis(mask: np.ndarray) -> ContourAxis:
    y, x = np.nonzero(mask)
    h, w = mask.shape
    if len(x):
        center = (float(x.min()) + float(x.max()) + 1) / 2
        top, bottom = float(y.min()), float(y.max() + 1)
    else:
        center, top, bottom = w / 2, 0.0, float(h)
    return ContourAxis((center, top), (center, max(top + 1, bottom)))


def mask_warnings(mask: np.ndarray) -> tuple[str, ...]:
    if not mask.any():
        return ("没有可测量的轮廓；请框选识别或手动补画。",)
    result = []
    if any(edge.any() for edge in (mask[0], mask[-1], mask[:, 0], mask[:, -1])):
        result.append("轮廓触及照片边界，可能截断；外缘与端部结果需复核。")
    return tuple(result)


def automatic_mask(
    rgba: np.ndarray,
    *,
    method: str = "auto",
    roi: tuple[int, int, int, int] | None = None,
    token: CancellationToken | None = None,
) -> tuple[np.ndarray, tuple[str, ...]]:
    """Suggest a native-size mask without adding a model dependency.

    A border-color distance works for a plain acquisition mat regardless of
    garment color. Dark/light Otsu modes and a crop are explicit alternatives.
    Largest component selection is a suggestion and is reported to the user.
    """
    _check(token)
    h, w = rgba.shape[:2]
    if h * w > MAX_PIXELS:
        raise ValueError("单图最多支持 3600 万像素。")
    if method not in {"auto", "dark", "light"}:
        raise ValueError("不支持的轮廓识别方法。")
    x0, y0, x1, y1 = (0, 0, w, h) if roi is None else roi
    x0, y0, x1, y1 = max(0, x0), max(0, y0), min(w, x1), min(h, y1)
    if x1 - x0 < 4 or y1 - y0 < 4:
        raise ValueError("框选区域过小，请包含完整物体和少量背景。")
    image = np.ascontiguousarray(rgba[y0:y1, x0:x1])
    alpha = image[:, :, 3]
    warnings = []
    if alpha.min() < 128 and alpha.max() >= 128:
        binary = (alpha >= 128).astype(np.uint8)
    else:
        rgb = image[:, :, :3]
        if method == "auto":
            lab = cv2.cvtColor(rgb, cv2.COLOR_RGB2LAB)
            border = np.concatenate((lab[0, ::2], lab[-1, ::2], lab[::2, 0], lab[::2, -1]))
            bg = np.median(border, axis=0).astype(np.uint8)
            diff = cv2.absdiff(lab, tuple(map(int, bg)))
            score = np.max(diff, axis=2)
            spread = np.median(np.max(np.abs(border.astype(float) - bg), axis=1))
            if spread > 18:
                warnings.append("背景不够均匀；建议框选物体或使用深色／浅色识别后修正。")
        else:
            score = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
        if int(score.max()) - int(score.min()) < 4:
            return np.zeros((h, w), dtype=bool), ("图像对比不足，请手动圈画轮廓。",)
        threshold_mode = cv2.THRESH_BINARY_INV if method == "dark" else cv2.THRESH_BINARY
        _, binary = cv2.threshold(score, 0, 1, threshold_mode | cv2.THRESH_OTSU)
    _check(token)
    count, labels, stats, _ = cv2.connectedComponentsWithStats(binary, connectivity=8)
    if count < 2:
        return np.zeros((h, w), dtype=bool), ("未识别到物体，请手动圈画轮廓。",)
    index = 1 + int(np.argmax(stats[1:, cv2.CC_STAT_AREA]))
    largest = labels == index
    other_area = int(binary.sum()) - int(largest.sum())
    if other_area > max(25, int(largest.sum()) * .01):
        warnings.append("自动保留最大连通物体；请检查是否遗漏断开的部分。")
    # Fill enclosed holes only. The gap between trouser legs remains exterior.
    flood = np.pad(largest.astype(np.uint8), 1)
    cv2.floodFill(flood, None, (0, 0), 2)
    filled = flood[1:-1, 1:-1] != 2
    if np.any(filled & ~largest):
        warnings.append("自动补齐了封闭内部空白；若为真实开孔，请手动剔除。")
    largest = filled
    if roi is not None and any(edge.any() for edge in (largest[0], largest[-1], largest[:, 0], largest[:, -1])):
        warnings.append("轮廓触及框选边界，可能被截断；请扩大框选范围后复核。")
    output = np.zeros((h, w), dtype=bool)
    output[y0:y1, x0:x1] = largest
    _check(token)
    warnings.extend(mask_warnings(output))
    if output.mean() > .9:
        warnings.append("识别范围接近整张照片，可能选中了背景，请复核。")
    return output, tuple(dict.fromkeys(warnings))


def frame_from_rgba(
    rgba: np.ndarray, label: str, *, source_path: str = "",
    mm_per_pixel: float | None = None, token: CancellationToken | None = None,
) -> ContourFrame:
    if rgba.ndim != 3 or rgba.shape[2] != 4 or min(rgba.shape[:2]) < 2:
        raise ValueError("输入需要一张有效的二维 RGBA 照片。")
    if rgba.shape[0] * rgba.shape[1] > MAX_PIXELS:
        raise ValueError("单图最多支持 3600 万像素。")
    _check(token)
    pixels = np.array(rgba, dtype=np.uint8, order="C", copy=True)
    mask, warnings = automatic_mask(pixels, token=token)
    return ContourFrame(label, pixels, mask, default_axis(mask), mm_per_pixel, warnings, source_path)


def load_frame(path: str | Path, *, token: CancellationToken | None = None) -> ContourFrame:
    with Image.open(path) as image:
        if getattr(image, "n_frames", 1) != 1:
            raise ValueError("仅支持单张二维照片；请先从图像堆栈导出所需的一帧。")
        if image.width * image.height > MAX_PIXELS:
            raise ValueError("单图最多支持 3600 万像素，请使用完整物体的合适分辨率照片。")
        image = ImageOps.exif_transpose(image)
        rgba = np.asarray(image.convert("RGBA"))
    return frame_from_rgba(rgba, Path(path).name, source_path=str(path), token=token)


def edit_mask(frame: ContourFrame, points, *, add: bool, radius: float = 0, polygon: bool = False) -> ContourFrame:
    points = np.asarray(points, dtype=float)
    if points.ndim != 2 or points.shape[1] != 2 or not np.isfinite(points).all():
        raise ValueError("修正坐标无效。")
    if polygon and len(points) < 3:
        raise ValueError("圈画轮廓至少需要三个点。")
    if not len(points):
        return frame
    result = frame.mask.astype(np.uint8)
    # Qt image coordinates describe pixel cells [x,x+1); clamp off-image drags
    # before converting to the OpenCV integer drawing coordinates.
    xy = np.rint(points - .5).clip(-100000, 100000).astype(np.int32)
    value = int(add)
    if polygon:
        cv2.fillPoly(result, [xy], value)
    else:
        radius = _positive(radius, "画笔半径")
        width = max(1, int(round(radius * 2)))
        for first, second in zip(xy, xy[1:]):
            cv2.line(result, tuple(first), tuple(second), value, width)
        for point in (xy[0], xy[-1]):
            cv2.circle(result, tuple(point), max(1, int(round(radius))), value, -1)
    mask = result.astype(bool)
    return replace(frame, mask=mask, warnings=mask_warnings(mask), edited=True)


def boundary_segments(mask: np.ndarray) -> np.ndarray:
    """Exact pixel-cell edges, including holes and disconnected components."""
    padded = np.pad(mask, 1)
    segments = []
    edge_count = 0
    for neighbor, start, end in (
        (padded[:-2, 1:-1], (0, 0), (1, 0)),
        (padded[2:, 1:-1], (0, 1), (1, 1)),
        (padded[1:-1, :-2], (0, 0), (0, 1)),
        (padded[1:-1, 2:], (1, 0), (1, 1)),
    ):
        y, x = np.nonzero(mask & ~neighbor)
        edge_count += len(x)
        if edge_count > 2_000_000:
            raise ValueError("轮廓碎片过多，请先剔除噪声或框选目标。")
        p = np.column_stack((x, y)).astype(np.float64)
        segments.append(np.stack((p + start, p + end), axis=1))
    output = np.concatenate(segments)
    return output


def _sections(segments: np.ndarray, heights: np.ndarray, step: float, token) -> tuple[tuple[tuple[float, float], ...], ...]:
    """Sweep only edges that intersect the fixed-height sample grid."""
    count = len(heights)
    first, second = segments[:, 0], segments[:, 1]
    dy = second[:, 1] - first[:, 1]
    low = np.minimum(first[:, 1], second[:, 1])
    high = np.maximum(first[:, 1], second[:, 1])
    # Half-open vertical intervals prevent duplicate crossings at vertices.
    starts = np.clip(np.ceil((low - heights[0]) / step - 1e-10), 0, count).astype(int)
    stops = np.clip(np.ceil((high - heights[0]) / step - 1e-10), 0, count).astype(int)
    hits: list[list[float]] = [[] for _ in heights]
    active = np.flatnonzero((stops > starts) & (np.abs(dy) > 1e-12))
    for offset, edge in enumerate(active):
        if offset % 2048 == 0:
            _check(token)
        for row in range(starts[edge], stops[edge]):
            fraction = (heights[row] - first[edge, 1]) / dy[edge]
            hits[row].append(float(first[edge, 0] + fraction * (second[edge, 0] - first[edge, 0])))
    rows = []
    for values in hits:
        values.sort()
        if len(values) % 2:
            raise ValueError("轮廓截线存在未闭合交点，请检查掩膜。")
        intervals = []
        for left, right in zip(values[::2], values[1::2]):
            if right - left <= 1e-8:
                continue
            if intervals and abs(left - intervals[-1][1]) < 1e-8:
                intervals[-1] = (intervals[-1][0], right)
            else:
                intervals.append((left, right))
        rows.append(tuple(intervals))
    return tuple(rows)


def compare_contours(before: ContourFrame, after: ContourFrame, step: float, *, token=None) -> ContourComparison:
    step = _positive(step, "采样间距")
    calibrated = (before.mm_per_pixel is not None, after.mm_per_pixel is not None)
    if calibrated[0] != calibrated[1]:
        raise ValueError("请为两张图片设置标定；同一拍摄比例时可沿用处理前标定。")
    if not any(calibrated) and before.mask.shape != after.mask.shape:
        raise ValueError("未标定且图片尺寸不同，不能直接比较像素；请分别标定。")
    if not before.mask.any() or not after.mask.any():
        raise ValueError("两张图片都需要有效轮廓。可自动识别，也可手动补画。")
    worlds, bounds = [], []
    for frame in (before, after):
        _check(token)
        world = frame.axis.to_world(boundary_segments(frame.mask), frame.scale)
        flat = world.reshape(-1, 2)
        bounds.append(tuple(map(float, (*flat.min(axis=0), *flat.max(axis=0)))))
        worlds.append(world)
    low, high = min(b[1] for b in bounds), max(b[3] for b in bounds)
    start = math.ceil(low / step - .5 - 1e-10) * step + step / 2
    count = max(1, int(math.ceil((high - start) / step)))
    if count > MAX_SAMPLES:
        raise ValueError(f"采样超过 {MAX_SAMPLES} 行，请增大采样间距（建议至少 {(high-low)/MAX_SAMPLES:.4g}）。")
    heights = start + np.arange(count) * step
    if count == 1 and heights[0] >= high:
        heights[0] = (low + high) / 2
    profiles = [_sections(world, heights, step, token) for world in worlds]
    result = []

    def measures(intervals):
        if not intervals:
            return None, None, None
        left, right = intervals[0][0], intervals[-1][1]
        return (-left if left <= 0 else None, right if right >= 0 else None, right - left)

    def delta(a, b):
        return None if a is None or b is None else b - a

    for height, b, a in zip(heights, *profiles):
        bl, br, bw = measures(b)
        al, ar, aw = measures(a)
        status = "可比较" if b and a else ("仅处理前有轮廓" if b else ("仅处理后有轮廓" if a else "无轮廓"))
        if b and a and len(b) != len(a):
            status = "分段数改变；外缘可比较，内部边界需复核"
        boundary_changes = ()
        if b and a and len(b) == len(a):
            boundary_changes = tuple(
                tuple(abs(new) - abs(old) if old * new >= 0 else None for old, new in zip(first, second))
                for first, second in zip(b, a)
            )
        result.append(ContourSection(float(height), b, a, bl, al, br, ar, delta(bl, al), delta(br, ar), bw, aw, delta(bw, aw), status, boundary_changes))
    warnings = list(before.warnings + after.warnings)
    warnings.extend(mask_warnings(before.mask) + mask_warnings(after.mask))
    if not any(calibrated):
        warnings.insert(0, "未标定：结果为 px / px²；仅适用于相同拍摄比例，不能解释为毫米。")
    if not (before.axis_confirmed and after.axis_confirmed):
        warnings.append("中线为自动建议，尚未人工核对；第一点应是共同参考高度，第二点朝向下摆。")
    warnings.append("同高度外形比较包含摆放差异，不代表同一材料点的收缩或应变。")
    warnings.append("各段边界按同高度从左至右配对；分段数改变或越过中线时不推断对应。")
    return ContourComparison("mm" if all(calibrated) else "px", step, tuple(result), bounds[0], bounds[1], float(before.mask.sum()) * before.scale**2, float(after.mask.sum()) * after.scale**2, tuple(dict.fromkeys(warnings)))


def _atomic(path: str | Path, writer: Callable[[str], None]) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=f".{path.name}.", suffix=path.suffix, dir=path.parent)
    os.close(fd)
    try:
        writer(temporary)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def save_comparison(path, before: ContourFrame | None, after: ContourFrame | None, step: float, *, same_capture: bool = True) -> None:
    """Portable session: embedded source pixels, native masks and independent axes."""
    def write(temporary):
        manifest = {"schema": SCHEMA, "version": 1, "step": _positive(step, "采样间距"), "same_capture": bool(same_capture), "frames": {}}
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_STORED) as archive:
            for key, frame in (("before", before), ("after", after)):
                if frame is None:
                    continue
                manifest["frames"][key] = {"label": frame.label, "axis": asdict(frame.axis), "mm_per_pixel": frame.mm_per_pixel, "warnings": list(frame.warnings), "source_path": frame.source_path, "edited": frame.edited, "axis_confirmed": frame.axis_confirmed}
                for name, array in (("image", frame.rgba), ("mask", frame.mask.astype(np.uint8) * 255)):
                    stream = BytesIO()
                    Image.fromarray(array).save(stream, format="PNG")
                    archive.writestr(f"{key}/{name}.png", stream.getvalue())
            archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, allow_nan=False))
    _atomic(path, write)


def load_comparison(path) -> tuple[ContourFrame | None, ContourFrame | None, float, bool]:
    frames = []
    with zipfile.ZipFile(path) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        if manifest.get("schema") != SCHEMA or manifest.get("version") != 1:
            raise ValueError("不支持的轮廓对比文件版本。")
        for key in ("before", "after"):
            data = manifest["frames"].get(key)
            if data is None:
                frames.append(None)
                continue
            images = []
            for name, mode in (("image", "RGBA"), ("mask", "L")):
                with archive.open(f"{key}/{name}.png") as member, Image.open(member) as image:
                    if image.width * image.height > MAX_PIXELS:
                        raise ValueError("对比文件中的图像超过像素上限。")
                    images.append(np.array(image.convert(mode)))
            axis = ContourAxis(**data["axis"])
            frames.append(ContourFrame(str(data["label"]), images[0], images[1] > 127, axis, data.get("mm_per_pixel"), tuple(map(str, data.get("warnings", []))), str(data.get("source_path", "")), bool(data.get("edited", False)), bool(data.get("axis_confirmed", False))))
    return frames[0], frames[1], _positive(manifest["step"], "采样间距"), bool(manifest.get("same_capture", True))


def export_comparison_excel(path, before: ContourFrame, after: ContourFrame, result: ContourComparison) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总与口径"
    for row in (("项目", "数值", "单位 / 说明"), ("处理前", before.label, ""), ("处理后", after.label, ""), ("比较口径", "相同参考高度的外形", "正值向外，负值向内；不自动缩放衣物"), ("采样间距", result.step, result.unit)):
        summary.append(row)
    for label, value in result.summary.items():
        unit = "%" if "%" in label else (result.unit + "²" if "面积" in label else result.unit)
        summary.append((label, value, unit))
    for stage, frame in (("处理前", before), ("处理后", after)):
        summary.append((f"{stage}中线起点 (px)", str(frame.axis.origin), "共同参考高度；第二点仅决定方向"))
        summary.append((f"{stage}中线方向点 (px)", str(frame.axis.direction_point), "不使用中线长度归一化"))
        summary.append((f"{stage}像素比例", frame.mm_per_pixel, "mm/px；空白表示未标定"))
        summary.append((f"{stage}中线状态", "人工设置" if frame.axis_confirmed else "自动建议，待核对", "第一点固定参考高度，第二点只决定方向"))
    for warning in result.warnings:
        summary.append(("说明", warning, ""))
    sheet = workbook.create_sheet("逐高度外缘")
    sheet.append([f"参考高度 ({result.unit})", "前左距", "后左距", "左变化", "前右距", "后右距", "右变化", "前外缘跨度", "后外缘跨度", "跨度变化", "状态", "前分段 [左,右]", "后分段 [左,右]", "各段边界距中线变化 [左,右]"])
    for row in result.sections:
        sheet.append([row.height, row.left_before, row.left_after, row.left_change, row.right_before, row.right_after, row.right_change, row.span_before, row.span_after, row.span_change, row.status, json.dumps(row.before_intervals, allow_nan=False), json.dumps(row.after_intervals, allow_nan=False), json.dumps(row.boundary_changes, allow_nan=False)])
    for target in workbook:
        target.freeze_panes = "A2"
        target.auto_filter.ref = target.dimensions
        for cell in target[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="245A70")
        for row in target.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (float, int)):
                    cell.number_format = "0.000"
                elif isinstance(cell.value, str):
                    # Labels are data, including filenames beginning with '='.
                    cell.data_type = "s"
        for index in range(1, target.max_column + 1):
            target.column_dimensions[get_column_letter(index)].width = 22 if index < 11 else 44
    summary.column_dimensions["B"].width = 80
    try:
        _atomic(path, lambda temporary: workbook.save(temporary))
    finally:
        workbook.close()
